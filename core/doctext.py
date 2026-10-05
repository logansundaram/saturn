"""Text out of the document formats people actually have — PDF, Word, Excel — for the direct
readers (`read_file`, `@file` attachments) and the knowledge-base loaders (stores/rag.py).

A leaf: imports nothing project-side, and the format libraries (pypdf, python-docx) load lazily
inside the one function that needs them, so importing this costs a plain launch nothing. Excel is
read with the standard library (an .xlsx is zipped XML), so it needs no extra package.

`extract(path)` is the direct readers' entry point: the document's text for a PDF / .docx /
.xlsx, else None (the caller reads the file as plain text — HTML and CSV are text already).
"""

import csv
import io
import re
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree


def extract(path) -> "str | None":
    """The readable text of a PDF / .docx / .xlsx, or None for any other suffix. A library or
    parse failure raises — the caller's tool round reports it as the error it is."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages = pdf_pages(path)
        text = "\n\n".join(f"--- page {n} ---\n{t}" for n, t in enumerate(pages, 1) if t.strip())
        return text or (f"[{path.name}: {len(pages)} page(s), no extractable text — "
                        "probably a scanned image]")
    if suffix == ".docx":
        return docx_to_text(path) or f"[{path.name}: no text in the document]"
    if suffix == ".xlsx":
        return xlsx_to_text(path) or f"[{path.name}: no cell values in the workbook]"
    return None


# ── PDF ─────────────────────────────────────────────────────────────────────────────────────────
# pypdf's raw extraction keeps page furniture (running headers/footers, page numbers) and hard
# line breaks mid-word. Both poison retrieval: furniture lines repeat into many chunks and match
# everything weakly; split words match nothing. Cleaned BEFORE chunking so the vectors only ever
# see content.
_DIGITS_RE = re.compile(r"\d+")


def pdf_pages(path) -> list[str]:
    """One cleaned text per page (furniture stripped, hyphenation repaired); a page with no text
    layer is an empty string, so indices stay page numbers − 1."""
    import logging

    import pypdf  # lazy (~56ms): only a PDF read needs it, never a plain launch

    # pypdf warns once per odd object in a real-world PDF; unhandled, those lines print to
    # stderr on top of the live TUI. A file it cannot read still raises.
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    reader = pypdf.PdfReader(str(path))
    raw_pages = [page.extract_text() or "" for page in reader.pages]
    return [normalize_pdf_text(t) for t in strip_repeated_furniture(raw_pages)]


def _furniture_key(line: str) -> str:
    """Normalize a candidate header/footer line for cross-page comparison: digits collapse so
    'Page 3 of 12' and 'Page 4 of 12' read as the same repeated line."""
    return _DIGITS_RE.sub("#", line.strip().lower())


def strip_repeated_furniture(pages: list[str], edge: int = 2, ratio: float = 0.6) -> list[str]:
    """Drop running headers/footers: a (digit-normalized) line that opens or closes most pages is
    page furniture, not content. Only the `edge` outermost lines of each page are candidates, so a
    sentence legitimately repeated mid-page is never touched. No-op for short documents (<3 pages),
    where 'repeated across pages' isn't meaningful."""
    if len(pages) < 3:
        return pages
    heads: Counter = Counter()
    feet: Counter = Counter()
    split_pages = [p.splitlines() for p in pages]
    edges = []  # per page: (head line indices, foot line indices) — the furniture candidates
    for lines in split_pages:
        content_idx = [i for i, ln in enumerate(lines) if ln.strip()]
        head_idx, foot_idx = content_idx[:edge], content_idx[-edge:]
        edges.append((head_idx, foot_idx))
        for i in head_idx:
            heads[_furniture_key(lines[i])] += 1
        for i in foot_idx:
            feet[_furniture_key(lines[i])] += 1
    threshold = max(3, int(len(pages) * ratio))
    head_junk = {k for k, c in heads.items() if c >= threshold}
    foot_junk = {k for k, c in feet.items() if c >= threshold}
    if not head_junk and not foot_junk:
        return pages

    cleaned = []
    for lines, (head_idx, foot_idx) in zip(split_pages, edges):
        drop = {i for i in head_idx if _furniture_key(lines[i]) in head_junk}
        drop |= {i for i in foot_idx if _furniture_key(lines[i]) in foot_junk}
        cleaned.append("\n".join(ln for i, ln in enumerate(lines) if i not in drop))
    return cleaned


def normalize_pdf_text(text: str) -> str:
    """Repair extraction artifacts: rejoin words hyphenated across line breaks, strip trailing
    whitespace, collapse blank-line runs."""
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ── Word ────────────────────────────────────────────────────────────────────────────────────────
def docx_to_text(path) -> str:
    """Text from a .docx: paragraphs in order, plus table cells row by row (tab-joined) — the two
    places Word documents keep their prose. Needs python-docx; a missing package raises a clear,
    actionable error."""
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise RuntimeError(
            "reading .docx needs the python-docx package — run `pip install python-docx`"
        ) from exc
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append("\t".join(cells))
    return "\n\n".join(parts)


# ── Excel ───────────────────────────────────────────────────────────────────────────────────────
# An .xlsx is a zip of SpreadsheetML: xl/workbook.xml names the sheets, its .rels maps each to a
# worksheet part, xl/sharedStrings.xml holds the text cells by index. Values come out as stored —
# a date is its serial number and a formula its cached result, which is what a reader needs.
_XL_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_XL_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_XLSX_MAX_ROWS = 2000                 # per sheet; the observation clamp would cut the rest anyway
_XLSX_MAX_PART_BYTES = 50_000_000     # uncompressed; refuse a zip bomb before parsing it
_CELL_REF_RE = re.compile(r"([A-Z]+)")


def _xml(zf: zipfile.ZipFile, name: str):
    if zf.getinfo(name).file_size > _XLSX_MAX_PART_BYTES:
        raise ValueError(f"{name} is too large to read ({zf.getinfo(name).file_size:,} bytes)")
    return ElementTree.fromstring(zf.read(name))


def _column_index(ref: str) -> int:
    """'A1' → 0, 'AB7' → 27."""
    m = _CELL_REF_RE.match(ref or "")
    if not m:
        return -1
    n = 0
    for ch in m.group(1):
        n = n * 26 + (ord(ch) - ord("A") + 1)
    return n - 1


def _text_of(node) -> str:
    """All <t> text under a shared-string or inline-string node (rich text is split in runs)."""
    return "".join(t.text or "" for t in node.iter(f"{_XL_MAIN}t"))


def _sheet_rows(root, shared: list[str]) -> "tuple[list[list[str]], bool]":
    rows: list[list[str]] = []
    capped = False
    for row in root.iter(f"{_XL_MAIN}row"):
        if len(rows) >= _XLSX_MAX_ROWS:
            capped = True
            break
        values: list[str] = []
        for cell in row.findall(f"{_XL_MAIN}c"):
            kind = cell.get("t", "n")
            if kind == "inlineStr":
                value = _text_of(cell)
            else:
                v = cell.find(f"{_XL_MAIN}v")
                value = v.text if v is not None and v.text is not None else ""
                if kind == "s" and value:
                    value = shared[int(value)] if int(value) < len(shared) else ""
                elif kind == "b" and value:
                    value = "TRUE" if value == "1" else "FALSE"
            col = _column_index(cell.get("r", ""))
            if col < 0:
                col = len(values)
            values.extend([""] * (col + 1 - len(values)))
            values[col] = value
        while values and not values[-1]:
            values.pop()
        rows.append(values)
    while rows and not rows[-1]:
        rows.pop()
    return rows, capped


def xlsx_to_text(path) -> str:
    """Every sheet of an .xlsx as CSV under a `## Sheet: <name>` heading, in workbook order."""
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        shared = []
        if "xl/sharedStrings.xml" in names:
            shared = [_text_of(si) for si in _xml(zf, "xl/sharedStrings.xml").iter(f"{_XL_MAIN}si")]
        targets = {}
        if "xl/_rels/workbook.xml.rels" in names:
            for rel in _xml(zf, "xl/_rels/workbook.xml.rels").iter(f"{_PKG_REL}Relationship"):
                target = rel.get("Target", "")
                targets[rel.get("Id")] = (target.lstrip("/") if target.startswith("/")
                                          else "xl/" + target)
        parts = []
        for sheet in _xml(zf, "xl/workbook.xml").iter(f"{_XL_MAIN}sheet"):
            part = targets.get(sheet.get(f"{_XL_REL}id"))
            if part not in names:
                continue
            rows, capped = _sheet_rows(_xml(zf, part), shared)
            if not rows:
                continue
            buf = io.StringIO()
            csv.writer(buf, lineterminator="\n").writerows(rows)
            body = buf.getvalue().rstrip("\n")
            if capped:
                body += f"\n… stopped at {_XLSX_MAX_ROWS} rows"
            parts.append(f"## Sheet: {sheet.get('name', '?')}\n{body}")
    return "\n\n".join(parts)
